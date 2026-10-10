#!/usr/bin/env python3
"""Generate H3 placebo plan: matched-length neutral control.

Adds a placebo frame that matches length/structure of treatment frames
but without normative content, to separate generic instruction salience
from normative semantic framing.

Design: 4 models × 27 profiles × 2 burden × 5 reps = 1,080 API calls
"""

import csv
import openpyxl
from pathlib import Path

PROJECT = Path(__file__).resolve().parent
INPUT_XLSX = PROJECT / "data" / "llm_experiment_inputs.xlsx"

# Load data
wb = openpyxl.load_workbook(INPUT_XLSX, read_only=True)

# Models (4-model subset)
ws = wb['Models']
mheaders = [str(h) for h in next(ws.iter_rows(values_only=True))]
models = [dict(zip(mheaders, row)) for row in ws.iter_rows(values_only=True)]
# Models (4-model subset - MUST match original H3)
# Note: These models may not all be in current inputs.xlsx
H3_MODELS = ['GPT-4.1', 'Llama 3.1 70B Instruct', 'Mistral Large', 'Qwen3.7 Plus']

# Try to find matching models in inputs
ws = wb['Models']
mheaders = [str(h) for h in next(ws.iter_rows(values_only=True))]
models = [dict(zip(mheaders, row)) for row in ws.iter_rows(values_only=True)]

h3_models = []
for target in H3_MODELS:
    matches = [m for m in models if m['model'] == target]
    if matches:
        h3_models.extend(matches)
    else:
        # Try partial match
        matches = [m for m in models if target.lower() in m['model'].lower()]
        if matches:
            h3_models.extend(matches)
            print(f"  Matched {target} -> {matches[0]['model']}")
        else:
            print(f"  Warning: {target} not found in inputs.xlsx")

if len(h3_models) < 4:
    print(f"\nWarning: Only {len(h3_models)}/4 models found. Placebo will use available models.")
    print("For exact H3 replication, ensure all 4 models are in inputs.xlsx:")
    print("  - GPT-4.1")
    print("  - Llama 3.1 70B Instruct")
    print("  - Mistral Large")
    print("  - Qwen3.7 Plus")
print(f"H3 models: {[m['model'] for m in h3_models]}")

# Profiles
ws = wb['Profiles']
prows = list(ws.iter_rows(values_only=True))
pheaders = [str(h) for h in prows[0]]
profiles = [dict(zip(pheaders, row)) for row in prows[1:]]
print(f"Profiles: {len(profiles)}")

# Burden Scenarios
ws = wb['Burden Scenarios']
srows = list(ws.iter_rows(values_only=True))
bheaders = [str(h) for h in srows[0]]
scenarios = [dict(zip(bheaders, row)) for row in srows[1:]]
print(f"Scenarios: {len(scenarios)}")

wb.close()

# ── Frames: 5 conditions ──────────────────────────────────────────

FRAMES = {
    "neutral": (
        "You are an evidence-based social scientist. Predict objectively. "
        "Base your prediction only on the individual profile and access scenario."
    ),
    "placebo": (
        "You are an evidence-based social scientist. When predicting behavior, "
        "give special weight to information completeness, response clarity, "
        "and thoroughness of available documentation. Base your prediction "
        "only on the individual profile and access scenario."
    ),
    "autonomy": (
        "You are an evidence-based social scientist. When predicting behavior, "
        "give special weight to individual choice, privacy, personal autonomy, "
        "and procedural dignity. Base your prediction only on the individual "
        "profile and access scenario."
    ),
    "collective": (
        "You are an evidence-based social scientist. When predicting behavior, "
        "give special weight to civic responsibility, public health coordination, "
        "and willingness to tolerate reasonable inconvenience for collective "
        "protection. Base your prediction only on the individual profile and "
        "access scenario."
    ),
    "equity": (
        "You are an evidence-based social scientist. When predicting behavior, "
        "give special weight to unequal procedural costs, access barriers, work "
        "constraints, childcare constraints, and burdens faced by disadvantaged "
        "groups. Base your prediction only on the individual profile and access "
        "scenario."
    ),
}

# Verify length matching
print("\nFrame lengths (characters):")
for name, text in FRAMES.items():
    print(f"  {name:12}: {len(text)}")

USER_PROMPT_TEMPLATE = """Person profile:
{profile_text}

Scenario:
{burden_text}

Based on the person described above, what is their willingness (0-100) to proceed with vaccination given this registration scenario?"""

# ── Generate plan ─────────────────────────────────────────────────

REPETITIONS = 5

rows = []
run_id = 1

for model_info in h3_models:
    for frame_name in ["placebo"]:  # Only placebo (new condition)
        for profile in profiles:
            for scenario in scenarios:
                for rep in range(1, REPETITIONS + 1):
                    user_prompt = USER_PROMPT_TEMPLATE.format(
                        profile_text=profile['profile_text'],
                        burden_text=scenario['scenario_text'],
                    )
                    rows.append({
                        'run_id': f"H3P_{run_id:05d}",
                        'model': model_info['model'],
                        'provider': model_info['provider'],
                        'deployment': model_info['deployment'],
                        'api_model': model_info['api_model'],
                        'access_channel': model_info.get('access_channel', ''),
                        'frame': frame_name,
                        'system_prompt': FRAMES[frame_name],
                        'profile_id': profile['profile_id'],
                        'profile_text': profile['profile_text'],
                        'burden_level': scenario['burden_level'],
                        'scenario_text': scenario['scenario_text'],
                        'repetition': rep,
                        'user_prompt': user_prompt,
                    })
                    run_id += 1

print(f"\nTotal runs: {len(rows)} (expected: 1,080)")

# Save
fields = ['run_id', 'model', 'provider', 'deployment', 'api_model', 'access_channel',
          'frame', 'system_prompt', 'profile_id', 'profile_text',
          'burden_level', 'scenario_text', 'repetition', 'user_prompt']

output_file = PROJECT / "data" / "h3_placebo_plan.csv"
with open(output_file, "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fields)
    w.writeheader()
    w.writerows(rows)

print(f"Saved: {output_file}")
